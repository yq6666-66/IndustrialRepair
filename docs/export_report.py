# -*- coding: utf-8 -*-
"""实训报告草稿.md → 按任务书格式排版的 Word（docs/实训报告_工业维修智能化零件分配平台.docx）

格式（任务书书写格式·2026-10 版）：
  一级标题：黑体·二号(22pt)加粗
  二级标题：黑体·三号(16pt)加粗
  正文：    宋体·小四(12pt)，行距固定值 22 磅，首行缩进 2 字符
  表格：    宋体·五号(10.5pt)，表头加粗
  封面：    实训报告封面（课程名称/课题名称/专业/班级/学号/姓名/指导教师/日期/校名）
  目录：    TOC 域；存在 toc_pages.json 时预填带页码目录（两遍法刷新见 refresh_toc_pages.py），Word 中仍可 F9 重建
  页码：    正文另起节，页脚居中，页码从 1 开始（装订要求：起始页码从正文开始）
  插图：    草稿中 ![图注](相对路径) 渲染为居中图片 + 图注（宋体五号居中）
用法：python docs/export_report.py
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "实训报告草稿.md"
OUT = ROOT / "实训报告_工业维修智能化零件分配平台.docx"
TOC_CACHE = ROOT / "toc_pages.json"

HEI, SONG = "黑体", "宋体"
BLACK = RGBColor(0, 0, 0)


def set_font(run, name=SONG, size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color      # 标题样式自带蓝色，强制黑字
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def fixed_line_spacing(par, pts=22):
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    par.paragraph_format.line_spacing = Pt(pts)


def add_toc(doc, entries=None):
    """插入 TOC 域。entries（toc_pages.json）非空时把目录与页码预填为域结果：
    Word 打开即见完整目录，仍可 Ctrl+A → F9 刷新。"""
    first = doc.add_paragraph()
    r = first.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r._r.append(fld); r._r.append(instr); r._r.append(sep)

    if not entries:
        tip = OxmlElement("w:t")
        tip.text = "（目录域：请在 Word 中右键→更新域，生成带页码目录）"
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        r._r.append(tip); r._r.append(end)
        return

    last = first
    for k, en in enumerate(entries):
        p = first if k == 0 else doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        fixed_line_spacing(p, 22)
        if en.get("level") == 2:
            p.paragraph_format.left_indent = Pt(24)   # 二级条目缩进 2 字符
        set_font(p.add_run(f'{en["text"]}\t{en["page"]}'), SONG, 12)
        last = p
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    last.add_run()._r.append(end)


def add_page_field(par):
    """在段落中插入 PAGE 域（页脚页码，宋体五号，居中由调用方设置）。"""
    r = par.add_run()
    set_font(r, SONG, 10.5)
    for kind in ("begin", "instr", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = "PAGE"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def setup_body_pagination(doc):
    """正文另起一节：页脚居中页码，且页码从 1 重新起算。"""
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec = doc.sections[-1]
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sec._sectPr.append(pg)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
IMG = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")


def add_body(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Pt(24)  # 小四 2 字符
    fixed_line_spacing(par)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1]); run.font.name = "Consolas"
            run.font.size = Pt(10.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
        else:
            set_font(par.add_run(part))


def add_table(doc, rows):
    cells = [ [c.strip() for c in r.strip().strip("|").split("|")] for r in rows ]
    ncol = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncol)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            run = par.add_run(re.sub(r"\*\*|`", "", txt))
            set_font(run, size=10.5, bold=(i == 0))
    # 表后空段
    sp = doc.add_paragraph(); fixed_line_spacing(sp, 12)


def main():
    if not SRC.exists():
        sys.exit(f"未找到 {SRC}")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Cm(3.0))
    sec.top_margin = sec.bottom_margin = Cm(2.5)

    # ---- 封面（按任务书"实训报告"封面样式）----
    def cover(text, size, bold=True, before=0, after=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        set_font(p.add_run(text), HEI, size, bold)

    for _ in range(4):
        cover("", 16, False)
    cover("实　训　报　告", 26, True, after=36)
    cover("课程名称：软件工程综合实训", 14, False, after=10)
    cover("课题名称：工业维修智能化零件分配平台", 14, False, after=10)
    cover("专　　业：软件工程", 14, False, after=10)
    cover("班　　级：软件工程 23 级（E511/E512）", 14, False, after=10)
    cover("学　　号：________________", 14, False, after=10)
    cover("姓　　名：________________", 14, False, after=10)
    cover("指导教师：谢雅、何潇、余新宇、刘明昕", 14, False, after=10)
    cover("", 16, False)
    cover("2026 年 10 月 14 日", 14, False)
    for _ in range(8):
        cover("", 16, False)
    cover("湖南工程学院", 18, False)
    doc.add_page_break()

    # ---- 目录 ----
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(h.add_run("目　　录"), HEI, 16)
    entries = None
    if TOC_CACHE.exists():
        try:
            entries = json.loads(TOC_CACHE.read_text(encoding="utf-8"))
        except Exception as e:  # 缓存损坏时退回空目录域
            print(f"（toc_pages.json 解析失败，改为空目录域：{e}）")
    add_toc(doc, entries)

    # ---- 正文另起一节：页码从 1 开始，页脚居中 ----
    setup_body_pagination(doc)

    # ---- 正文（解析 Markdown）----
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1; continue
        if line.startswith(">"):          # 草稿附注（封面信息/排版提醒）不进入成稿
            i += 1; continue
        if line.startswith("|"):          # 表格
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:\-|]+\|$", lines[i].strip()):
                    rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        m = IMG.match(line)
        if m:                             # 插图：![图注](相对路径)
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(6)
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(ROOT / m.group(2)), width=Cm(14))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            fixed_line_spacing(cap, 18)
            set_font(cap.add_run(m.group(1)), SONG, 10.5)
            i += 1; continue
        if line.startswith("# "):
            i += 1; continue              # 文档主标题（封面已含，正文不再重复）
        if line.startswith("## "):
            # 用内置标题样式（目录域 F9 重建时按样式收集，防止刷新为空）
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            set_font(p.add_run(line[3:].strip()), HEI, 22, True, color=BLACK)  # 一级：黑体·二号(22pt)加粗
            i += 1; continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            set_font(p.add_run(line[4:].strip()), HEI, 16, True, color=BLACK)  # 二级：黑体·三号(16pt)加粗
            i += 1; continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            fixed_line_spacing(p, 22)
            set_font(p.add_run(line[5:].strip()), HEI, 12, True, color=BLACK)  # 三级：黑体·小四（任务书未强制）
            i += 1; continue
        if re.match(r"^[-*] ", line) or re.match(r"^\d+\. ", line):
            par = doc.add_paragraph(style="List Bullet" if line.startswith(("-", "*")) else "List Number")
            fixed_line_spacing(par)
            txt = re.sub(r"^([-*]|\d+\.)\s+", "", line)
            for part in INLINE.split(txt):
                if part:
                    set_font(par.add_run(part.strip("*`")))
            i += 1; continue
        add_body(doc, line)
        i += 1

    doc.save(OUT)
    print(f"√ 已生成：{OUT}")
    toc_state = f"已预填 {len(entries)} 条带页码目录" if entries else "目录域为空（无 toc_pages.json）"
    print(f"  {toc_state}；封面学号/姓名待填；页码从正文第 1 页起（页脚居中）。")


if __name__ == "__main__":
    main()
